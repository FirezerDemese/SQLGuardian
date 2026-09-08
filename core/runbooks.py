"""
SQLGuardian - Runbook ingestion and condition-scoped retrieval

The product is not detection. Detection is the commodity input. The product is
the distance between a condition firing and a junior engineer knowing what THIS
team does about it - which lives in a Confluence page nobody opens at 3am.

This module ingests that page, keeps it in one piece, and hands back the
specific section that covers the specific condition, with a citation.

Design decisions worth defending:

  Chunk by heading, never by token count. A retrieved step that has lost its
  surrounding procedure is worse than no step at all, because it reads like it
  is complete. Every chunk keeps its full heading path.

  Retrieval is scoped by condition code first, ranked second. A blind
  similarity search over the whole corpus can return the disaster-recovery
  runbook for a blocking incident and be confidently wrong. Sections are
  eligible only if they are mapped to the condition that fired.

  Mapping is auditable. Either the document says which conditions it covers
  (an explicit annotation) or a keyword rule in this file matched, and the
  rule that matched is recorded and shown in the UI. There is no hidden
  scoring model deciding what your team's procedure applies to.

  No embeddings, no vector store, no network call. Retrieval is a deterministic
  function of the corpus and the condition. It gives the same answer at 3am
  during an outage as it did in testing.
"""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from types import MappingProxyType
from typing import Iterable, Mapping, Optional, Sequence

from loguru import logger

from config.settings import settings
from core.conditions import Condition, ConditionCode


SUPPORTED_EXTENSIONS: Mapping[str, str] = MappingProxyType({
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
    ".html": "html",
    ".htm": "html",
    ".pdf": "pdf",
    ".docx": "docx",
})


class RunbookIngestError(ValueError):
    """The document could not be parsed into cited sections."""


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class RunbookSection:
    """One procedure, kept whole, with everything needed to cite it."""

    doc_id: str
    doc_title: str
    doc_filename: str
    heading: str
    heading_path: tuple[str, ...]
    anchor: str
    body: str
    source_ref: str                       # "lines 40-58" or "page 3"
    conditions: tuple[str, ...]           # ConditionCode values this section covers
    mapping_reason: str                   # why it is mapped - shown in the UI

    @property
    def citation(self) -> str:
        """What gets printed next to every rendered step."""
        path = " > ".join(self.heading_path) if self.heading_path else self.heading
        return f"{self.doc_title} - {path} ({self.doc_filename}#{self.anchor}, {self.source_ref})"

    def steps(self) -> tuple[str, ...]:
        """Numbered or bulleted lines in the body, in document order.

        Falls back to the whole body as one step when the section is prose -
        we do not invent structure the document does not have.
        """
        found = tuple(
            re.sub(r"^\s*(?:[-*•]|\d+[.)])\s+", "", line).strip()
            for line in self.body.splitlines()
            if re.match(r"^\s*(?:[-*•]|\d+[.)])\s+\S", line)
        )
        return found or ((self.body.strip(),) if self.body.strip() else ())


@dataclass(frozen=True)
class RunbookDocument:
    doc_id: str
    filename: str
    title: str
    format: str
    section_count: int
    mapped_section_count: int
    byte_size: int
    ingested_at: str
    conditions_covered: tuple[str, ...]


@dataclass(frozen=True)
class RetrievedSection:
    """A section returned for a specific condition, with why it was chosen."""

    section: RunbookSection
    score: float
    match_reason: str

    def to_dict(self) -> dict:
        return {
            "doc_id": self.section.doc_id,
            "doc_title": self.section.doc_title,
            "doc_filename": self.section.doc_filename,
            "heading": self.section.heading,
            "heading_path": list(self.section.heading_path),
            "anchor": self.section.anchor,
            "citation": self.section.citation,
            "source_ref": self.section.source_ref,
            "body": self.section.body,
            "steps": list(self.section.steps()),
            "conditions": list(self.section.conditions),
            "mapping_reason": self.section.mapping_reason,
            "score": round(self.score, 3),
            "match_reason": self.match_reason,
        }


# ---------------------------------------------------------------------------
# Condition mapping
# ---------------------------------------------------------------------------
# Explicit annotation wins. A team can write, anywhere inside a section:
#     Conditions: BLOCKING_CHAIN, LONG_RUNNING_REQUEST
# or, to keep it invisible in rendered Markdown / Confluence:
#     <!-- sqlguardian: BLOCKING_CHAIN -->

_ANNOTATION_RE = re.compile(
    r"(?:<!--\s*sqlguardian\s*:|(?:^|\n)\s*(?:sqlguardian-)?conditions?\s*:)\s*"
    r"([A-Z_,\s]+?)\s*(?:-->|\n|$)",
    re.IGNORECASE,
)

# Fallback rules. Deliberately narrow: a phrase here should be one a DBA would
# only write when documenting that specific condition. The matched phrase is
# recorded on the section and shown in the UI, so a bad rule is visible rather
# than silent.
CONDITION_KEYWORDS: Mapping[ConditionCode, tuple[str, ...]] = MappingProxyType({
    ConditionCode.BLOCKING_CHAIN: (
        "blocking chain", "head blocker", "blocked session", "lock contention",
        "blocking session", "sp_who2", "lck_m_",
    ),
    ConditionCode.AGENT_JOB_FAILURE: (
        "agent job", "job failure", "failed job", "sql agent", "sysjobhistory",
        "job step failed",
    ),
    ConditionCode.BACKUP_AGE_EXCEEDED: (
        "backup failure", "missed backup", "backup window", "last full backup",
        "rpo", "backup age", "restore test",
    ),
    ConditionCode.LOG_GROWTH: (
        "transaction log", "log growth", "log full", "log_reuse_wait",
        "ldf", "log file is full", "9002",
    ),
    ConditionCode.DISK_PRESSURE: (
        "disk space", "disk full", "volume full", "free space", "low disk",
        "drive is full",
    ),
    ConditionCode.WAIT_SPIKE: (
        "wait stats", "wait type", "pageiolatch", "writelog", "resource_semaphore",
        "sos_scheduler_yield", "cxpacket", "threadpool",
    ),
    ConditionCode.LONG_RUNNING_REQUEST: (
        "long running quer", "long-running quer", "runaway quer", "query timeout",
        "slow query", "long running request",
    ),
})


def _extract_annotated_conditions(text: str) -> tuple[tuple[str, ...], str]:
    """Return codes the document explicitly claims, plus the reason string."""
    codes: list[str] = []
    for match in _ANNOTATION_RE.finditer(text):
        for token in re.split(r"[,\s]+", match.group(1).strip()):
            token = token.upper()
            if not token:
                continue
            if token in ConditionCode.__members__:
                if token not in codes:
                    codes.append(token)
            else:
                logger.warning(f"Runbook annotation names an unknown condition: {token}")
    if codes:
        return tuple(codes), "declared in the document"
    return (), ""


def map_conditions(heading_path: Sequence[str], body: str) -> tuple[tuple[str, ...], str]:
    """Decide which conditions a section covers, and record why."""
    annotated, reason = _extract_annotated_conditions(f"{' '.join(heading_path)}\n{body}")
    if annotated:
        return annotated, reason

    heading_text = " ".join(heading_path).lower()
    body_text = body.lower()

    matched: list[str] = []
    reasons: list[str] = []
    for code, phrases in CONDITION_KEYWORDS.items():
        for phrase in phrases:
            if phrase in heading_text:
                matched.append(code.value)
                reasons.append(f"{code.value}: heading contains \"{phrase}\"")
                break
            if phrase in body_text:
                matched.append(code.value)
                reasons.append(f"{code.value}: body contains \"{phrase}\"")
                break
    return tuple(matched), "; ".join(reasons)


# ---------------------------------------------------------------------------
# Parsers - each returns (heading_path, body, source_ref) triples
# ---------------------------------------------------------------------------

ParsedChunk = tuple[tuple[str, ...], str, str]

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*\S)\s*#*\s*$")
_SETEXT_RE = re.compile(r"^\s*(=|-){3,}\s*$")
_NUMBERED_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)[.)]?\s+([A-Z][^\n]{2,80})$")


def _assemble(marks: Sequence[dict]) -> list[ParsedChunk]:
    """Maintain the H1>H2>H3 stack so every chunk keeps its full heading path.

    marks: [{'level': int, 'heading': str, 'body': str, 'ref': str}, ...]
    """
    chunks: list[ParsedChunk] = []
    stack: list[tuple[int, str]] = []
    for mark in marks:
        level, heading = mark["level"], mark["heading"]
        while stack and stack[-1][0] >= level:
            stack.pop()
        stack.append((level, heading))
        path = tuple(h for _, h in stack)
        body = mark["body"].strip()
        if body:
            chunks.append((path, body, mark["ref"]))
    return chunks


def _parse_markdown(text: str) -> list[ParsedChunk]:
    lines = text.splitlines()
    marks: list[dict] = []
    current: Optional[dict] = None
    buffer: list[str] = []
    start_line = 1

    def flush(end_line: int) -> None:
        nonlocal buffer
        if current is not None:
            current["body"] = "\n".join(buffer)
            current["ref"] = f"lines {start_line}-{end_line}"
            marks.append(current)
        buffer = []

    in_fence = False
    for index, line in enumerate(lines, start=1):
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
        heading_match = None if in_fence else _MD_HEADING_RE.match(line)
        if heading_match:
            flush(index - 1)
            current = {"level": len(heading_match.group(1)), "heading": heading_match.group(2).strip()}
            start_line = index
        elif not in_fence and index < len(lines) and _SETEXT_RE.match(lines[index]) and line.strip():
            flush(index - 1)
            level = 1 if lines[index].strip().startswith("=") else 2
            current = {"level": level, "heading": line.strip()}
            start_line = index
        else:
            if current is None and line.strip():
                current = {"level": 1, "heading": "(untitled section)"}
                start_line = index
            if current is not None:
                buffer.append(line)
    flush(len(lines))

    # Drop the setext underline that follows a heading line.
    for mark in marks:
        mark["body"] = "\n".join(
            l for l in mark["body"].splitlines() if not _SETEXT_RE.match(l)
        )
    return _assemble(marks)


def _looks_like_plain_heading(line: str, next_line: Optional[str]) -> Optional[int]:
    """Heading detection for formats with no markup: plain text and PDF.

    Three deterministic signals, documented in the README so a team knows how
    to write a .txt runbook that chunks correctly:
      - a setext underline (=== or ---) on the following line
      - a short ALL-CAPS line
      - a numbered section like "2." or "2.1" followed by a capitalised title
    """
    stripped = line.strip()
    if not stripped or len(stripped) > 90:
        return None
    if next_line is not None and _SETEXT_RE.match(next_line):
        return 1 if next_line.strip().startswith("=") else 2
    letters = [c for c in stripped if c.isalpha()]
    if letters and all(c.isupper() for c in letters) and len(stripped.split()) <= 10:
        return 1
    numbered = _NUMBERED_HEADING_RE.match(stripped)
    if numbered:
        return 1 + numbered.group(1).count(".")
    return None


def _parse_text(text: str, ref_prefix: str = "lines") -> list[ParsedChunk]:
    lines = text.splitlines()
    marks: list[dict] = []
    current: Optional[dict] = None
    buffer: list[str] = []
    start_line = 1

    def flush(end_line: int) -> None:
        nonlocal buffer
        if current is not None:
            current["body"] = "\n".join(
                l for l in buffer if not _SETEXT_RE.match(l)
            )
            current["ref"] = f"{ref_prefix} {start_line}-{end_line}"
            marks.append(current)
        buffer = []

    for index, line in enumerate(lines, start=1):
        next_line = lines[index] if index < len(lines) else None
        level = _looks_like_plain_heading(line, next_line)
        if level is not None:
            flush(index - 1)
            current = {"level": level, "heading": line.strip()}
            start_line = index
        else:
            if current is None and line.strip():
                current = {"level": 1, "heading": "(untitled section)"}
                start_line = index
            if current is not None:
                buffer.append(line)
    flush(len(lines))
    return _assemble(marks)


def _parse_html(text: str) -> list[ParsedChunk]:
    """Confluence and wiki exports. Uses the stdlib parser - no new dependency."""
    from html.parser import HTMLParser

    class Extractor(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.marks: list[dict] = []
            self.current: Optional[dict] = None
            self.buffer: list[str] = []
            self.heading_level: Optional[int] = None
            self.heading_text: list[str] = []
            self.skip_depth = 0
            self.list_item = False
            self.title: Optional[str] = None
            self.in_title = False

        def handle_starttag(self, tag: str, attrs: list) -> None:
            if tag in ("script", "style"):
                self.skip_depth += 1
            elif tag == "title":
                self.in_title = True
            elif re.fullmatch(r"h[1-6]", tag):
                self._flush()
                self.heading_level = int(tag[1])
                self.heading_text = []
            elif tag == "li":
                self.list_item = True
                self.buffer.append("\n- ")
            elif tag in ("p", "br", "div", "tr"):
                self.buffer.append("\n")

        def handle_endtag(self, tag: str) -> None:
            if tag in ("script", "style"):
                self.skip_depth = max(0, self.skip_depth - 1)
            elif tag == "title":
                self.in_title = False
            elif re.fullmatch(r"h[1-6]", tag) and self.heading_level is not None:
                self.current = {
                    "level": self.heading_level,
                    "heading": " ".join("".join(self.heading_text).split()),
                }
                self.heading_level = None
            elif tag == "li":
                self.list_item = False

        def handle_data(self, data: str) -> None:
            if self.skip_depth:
                return
            if self.in_title:
                self.title = (self.title or "") + data
            elif self.heading_level is not None:
                self.heading_text.append(data)
            elif self.current is not None:
                self.buffer.append(data)

        def handle_comment(self, data: str) -> None:
            # Annotations ride in HTML comments so they stay invisible in
            # Confluence: <!-- sqlguardian: BLOCKING_CHAIN -->
            if "sqlguardian" in data.lower() and self.current is not None:
                self.buffer.append(f"\n<!--{data}-->\n")

        def _flush(self) -> None:
            if self.current is not None:
                body = "".join(self.buffer)
                body = "\n".join(line.strip() for line in body.splitlines())
                body = re.sub(r"\n{3,}", "\n\n", body).strip()
                self.current["body"] = body
                self.current["ref"] = f"heading \"{self.current['heading']}\""
                self.marks.append(self.current)
            self.buffer = []
            self.current = None

        def close(self) -> None:  # type: ignore[override]
            super().close()
            self._flush()

    parser = Extractor()
    parser.feed(text)
    parser.close()
    return _assemble(parser.marks)


def _parse_pdf(raw: bytes) -> list[ParsedChunk]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise RunbookIngestError("PDF support needs the pypdf package.") from exc

    import io

    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception as exc:
        # A malformed upload is a rejected document, not a server error.
        raise RunbookIngestError(f"This file is not a readable PDF: {exc}") from exc

    chunks: list[ParsedChunk] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        for path, body, _ in _parse_text(text, ref_prefix="line"):
            chunks.append((path, body, f"page {page_number}"))
    if not chunks:
        raise RunbookIngestError(
            "No extractable text in this PDF. Scanned or image-only PDFs are not "
            "supported - there is no OCR step in this pipeline."
        )
    return chunks


def _parse_docx(raw: bytes) -> list[ParsedChunk]:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise RunbookIngestError("DOCX support needs the python-docx package.") from exc

    import io

    document = docx.Document(io.BytesIO(raw))
    marks: list[dict] = []
    current: Optional[dict] = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if current is not None:
            current["body"] = "\n".join(buffer).strip()
            current["ref"] = f"heading \"{current['heading']}\""
            marks.append(current)
        buffer = []

    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "") if paragraph.style else ""
        heading_match = re.fullmatch(r"Heading (\d)", style)
        text = paragraph.text.strip()
        if heading_match and text:
            flush()
            current = {"level": int(heading_match.group(1)), "heading": text}
        elif text:
            if current is None:
                current = {"level": 1, "heading": "(untitled section)"}
            prefix = "- " if style.startswith("List") else ""
            buffer.append(prefix + text)
    flush()
    return _assemble(marks)


def parse_document(filename: str, raw: bytes) -> list[ParsedChunk]:
    """Dispatch on extension. Raises RunbookIngestError on anything unsupported."""
    suffix = Path(filename).suffix.lower()
    fmt = SUPPORTED_EXTENSIONS.get(suffix)
    if fmt is None:
        raise RunbookIngestError(
            f"Unsupported runbook format '{suffix or filename}'. Supported: "
            + ", ".join(sorted(SUPPORTED_EXTENSIONS))
        )
    if fmt == "pdf":
        return _parse_pdf(raw)
    if fmt == "docx":
        return _parse_docx(raw)

    text = raw.decode("utf-8-sig", errors="replace")
    if fmt == "markdown":
        return _parse_markdown(text)
    if fmt == "html":
        return _parse_html(text)
    return _parse_text(text)


def document_format(filename: str) -> str:
    return SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower(), "unknown")


# ---------------------------------------------------------------------------
# Anchors
# ---------------------------------------------------------------------------

def _slug(text: str) -> str:
    normalised = unicodedata.normalize("NFKD", text)
    ascii_text = normalised.encode("ascii", "ignore").decode("ascii").lower()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_text).strip("-")
    return slug or "section"


def make_anchor(heading_path: Sequence[str], existing: Iterable[str]) -> str:
    """A stable, human-readable anchor.

    Derived from the heading path, so re-ingesting an edited document keeps the
    same anchor for an unchanged section and old citations stay resolvable.
    Collisions get a numeric suffix rather than a hash, so the anchor still
    reads like the heading it points at.
    """
    base = _slug("-".join(heading_path[-2:]) if len(heading_path) > 1 else heading_path[0])
    taken = set(existing)
    if base not in taken:
        return base
    counter = 2
    while f"{base}-{counter}" in taken:
        counter += 1
    return f"{base}-{counter}"


# ---------------------------------------------------------------------------
# Store
# ---------------------------------------------------------------------------

class RunbookStore:
    """Corpus on disk: one JSON index plus a copy of every original document.

    Deliberately a flat file. The corpus is a few dozen procedures, it has to
    be greppable by a human during an incident, and it must not add a database
    to the deployment.
    """

    def __init__(self, root: Optional[Path] = None) -> None:
        self.root = Path(root or settings.runbook_store_dir)
        self.index_path = self.root / "index.json"
        self.originals = self.root / "originals"
        self._documents: dict[str, RunbookDocument] = {}
        self._sections: list[RunbookSection] = []
        self._loaded = False

    # -- persistence -------------------------------------------------------

    def _ensure_loaded(self) -> None:
        if self._loaded:
            return
        self._loaded = True
        if not self.index_path.exists():
            return
        try:
            data = json.loads(self.index_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            logger.error(f"Runbook index unreadable ({exc}); starting empty.")
            return
        self._documents = {
            d["doc_id"]: RunbookDocument(
                **{**d, "conditions_covered": tuple(d.get("conditions_covered", ()))}
            )
            for d in data.get("documents", [])
        }
        self._sections = [
            RunbookSection(
                **{
                    **s,
                    "heading_path": tuple(s["heading_path"]),
                    "conditions": tuple(s["conditions"]),
                }
            )
            for s in data.get("sections", [])
        ]

    def _save(self) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        payload = {
            "documents": [asdict(d) for d in self._documents.values()],
            "sections": [asdict(s) for s in self._sections],
            "saved_at": datetime.now(timezone.utc).isoformat(),
        }
        self.index_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    # -- ingestion ---------------------------------------------------------

    def ingest(self, filename: str, raw: bytes, title: Optional[str] = None) -> RunbookDocument:
        """Parse, chunk, map to conditions, and store one runbook document.

        Re-ingesting the same filename replaces the previous version: teams edit
        runbooks, and two copies of a procedure in the corpus is how a stale
        step gets cited during an incident.
        """
        self._ensure_loaded()
        chunks = parse_document(filename, raw)
        if not chunks:
            raise RunbookIngestError(
                f"{filename} produced no sections. A runbook needs at least one "
                f"heading so retrieved steps keep their procedure."
            )

        doc_id = hashlib.sha256(filename.encode("utf-8")).hexdigest()[:12]
        doc_title = title or Path(filename).stem.replace("_", " ").replace("-", " ").strip()

        anchors: list[str] = []
        sections: list[RunbookSection] = []
        for heading_path, body, source_ref in chunks:
            anchor = make_anchor(heading_path, anchors)
            anchors.append(anchor)
            conditions, reason = map_conditions(heading_path, body)
            sections.append(RunbookSection(
                doc_id=doc_id,
                doc_title=doc_title,
                doc_filename=filename,
                heading=heading_path[-1],
                heading_path=heading_path,
                anchor=anchor,
                body=body,
                source_ref=source_ref,
                conditions=conditions,
                mapping_reason=reason or "no condition matched",
            ))

        covered = tuple(sorted({c for s in sections for c in s.conditions}))
        document = RunbookDocument(
            doc_id=doc_id,
            filename=filename,
            title=doc_title,
            format=document_format(filename),
            section_count=len(sections),
            mapped_section_count=sum(1 for s in sections if s.conditions),
            byte_size=len(raw),
            ingested_at=datetime.now(timezone.utc).isoformat(),
            conditions_covered=covered,
        )

        self._sections = [s for s in self._sections if s.doc_id != doc_id] + sections
        self._documents[doc_id] = document

        self.originals.mkdir(parents=True, exist_ok=True)
        (self.originals / f"{doc_id}{Path(filename).suffix.lower()}").write_bytes(raw)
        self._save()

        logger.info(
            f"Ingested runbook '{filename}' | format={document.format} "
            f"sections={document.section_count} mapped={document.mapped_section_count} "
            f"conditions={','.join(covered) or 'none'}"
        )
        return document

    def ingest_path(self, path: Path) -> RunbookDocument:
        path = Path(path)
        return self.ingest(path.name, path.read_bytes())

    def remove(self, doc_id: str) -> bool:
        self._ensure_loaded()
        if doc_id not in self._documents:
            return False
        suffix = Path(self._documents[doc_id].filename).suffix.lower()
        self._documents.pop(doc_id)
        self._sections = [s for s in self._sections if s.doc_id != doc_id]
        original = self.originals / f"{doc_id}{suffix}"
        if original.exists():
            original.unlink()
        self._save()
        return True

    # -- reads -------------------------------------------------------------

    def documents(self) -> tuple[RunbookDocument, ...]:
        self._ensure_loaded()
        return tuple(sorted(self._documents.values(), key=lambda d: d.title.lower()))

    def sections(self) -> tuple[RunbookSection, ...]:
        self._ensure_loaded()
        return tuple(self._sections)

    def sections_for(self, code: ConditionCode) -> tuple[RunbookSection, ...]:
        self._ensure_loaded()
        return tuple(s for s in self._sections if code.value in s.conditions)

    def covered_conditions(self) -> tuple[str, ...]:
        self._ensure_loaded()
        return tuple(sorted({c for s in self._sections for c in s.conditions}))

    def uncovered_conditions(self) -> tuple[str, ...]:
        covered = set(self.covered_conditions())
        return tuple(c.value for c in ConditionCode if c.value not in covered)

    # -- retrieval ---------------------------------------------------------

    def retrieve(
        self,
        condition: Condition,
        limit: int = 3,
    ) -> tuple[RetrievedSection, ...]:
        """Return the team's procedure for this condition, best match first.

        Scoped: only sections mapped to this condition code are eligible. An
        empty result is a real answer - it means the team has no documented
        procedure for this condition, and core/gap_report.py records it.
        """
        candidates = self.sections_for(condition.code)
        if not candidates:
            return ()

        terms = _retrieval_terms(condition)
        scored: list[RetrievedSection] = []
        for section in candidates:
            heading_text = " ".join(section.heading_path).lower()
            body_text = section.body.lower()
            hits: list[str] = []
            score = 1.0  # eligible at all == mapped to this condition
            for term in terms:
                if term in heading_text:
                    score += 2.0
                    hits.append(f"heading matches \"{term}\"")
                elif term in body_text:
                    score += 1.0
                    hits.append(f"body matches \"{term}\"")
            if section.mapping_reason == "declared in the document":
                score += 1.5
                hits.insert(0, "document declares this condition explicitly")
            scored.append(RetrievedSection(
                section=section,
                score=score,
                match_reason="; ".join(hits[:4]) or f"mapped to {condition.code.value}",
            ))

        scored.sort(key=lambda r: (-r.score, r.section.anchor))
        return tuple(scored[:limit])


def _retrieval_terms(condition: Condition) -> tuple[str, ...]:
    """Terms drawn from the live condition, used only to rank within scope.

    Ranking terms come from the facts the checks actually collected - the wait
    type that spiked, the database names involved - so the section that names
    the affected database wins over the generic one.
    """
    terms: list[str] = []
    for key in ("dominant_wait_type", "head_blocker_program"):
        value = condition.facts.get(key)
        if isinstance(value, str) and value.strip():
            terms.append(value.lower())
    for key in ("database_names", "failed_job_names", "volume_mount_points"):
        for value in condition.facts.get(key) or ():
            if isinstance(value, str) and value.strip():
                terms.append(value.lower())
    return tuple(dict.fromkeys(terms))


# Module-level store, mirroring db_manager. Tests construct their own.
runbook_store = RunbookStore()
